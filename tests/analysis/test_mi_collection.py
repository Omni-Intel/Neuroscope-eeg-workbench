import json

from docx import Document

from neuroscope_eeg.analysis.assessment import assess
from neuroscope_eeg.io.native_recording import write_json
from neuroscope_eeg.reports.generate import generate_report
from tests.analysis.test_assessment import guided_fixture


def test_four_class_acquisition_report_has_no_binary_results(tmp_path, monkeypatch):
    plan, root = guided_fixture(tmp_path, '运动想象', mi_mode='four_class_collection_v4')
    device = json.loads((root / 'device.json').read_text(encoding='utf-8'))
    device['channel_names'] = ['Cz', 'FCz']  # Acquisition must not require a binary C3/C4 montage.
    write_json(root / 'device.json', device)
    def forbidden(*args, **kwargs):
        raise AssertionError('The legacy binary decoder must not run for four-class collection')
    monkeypatch.setattr('neuroscope_eeg.analysis.assessment.mi_metrics', forbidden)
    path = generate_report(root)
    latest = json.loads((root / 'analysis/latest.json').read_text(encoding='utf-8'))
    report = json.loads((root / 'analysis' / latest['generation'] / 'assessment.json').read_text(encoding='utf-8'))
    task = report['task']
    assert task['status'] == 'pending' and task['collection_only']
    assert [r['condition'] for r in task['trials']] == list(plan.mi_labels) * 2
    assert all(r['data_complete'] for r in task['trials'])
    assert all(r['imagery_end_sample_exclusive'] > r['imagery_start_sample'] for r in task['trials'])
    assert all(r['completed'] == 2 for r in task['collection_counts'])
    assert not {'all_trials', 'erd', 'response_evidence'} & task.keys()
    assert report['general_features']['status'] == 'available'
    assert report['analysis_version'] == 'assessment-v1.5.0'
    doc = Document(path)
    text = '\n'.join(p.text for p in doc.paragraphs) + '\n' + '\n'.join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    for expected in ('双脚', '舌头', '待更改', '待接入', '基线与任务功率', '信号质量与电极接触', '通用信号特征'):
        assert expected in text
    assert '0.6 判左手' not in text and '名义随机水平 50%' not in text
    assert not (root / 'analysis' / latest['generation'] / 'figures/task-confusion.png').exists()


def test_incomplete_four_class_trial_keeps_label_and_issue(tmp_path):
    _, root = guided_fixture(tmp_path, '运动想象', mi_mode='four_class_collection_v4')
    events = [json.loads(line) for line in (root / 'events.jsonl').read_text(encoding='utf-8').splitlines()]
    events = [e for e in events if not (e['phase'] == 'trial_complete' and e['payload'].get('formal_trial_index') == 3)]
    (root / 'events.jsonl').write_text(''.join(json.dumps(e, ensure_ascii=False)+'\n' for e in events), encoding='utf-8')
    result = assess(root)['task']
    trial = next(r for r in result['trials'] if r['formal_trial_index'] == 3)
    assert trial['condition'] == '舌头' and not trial['completed'] and not trial['data_complete']
    assert trial['collection_issue'] == '阶段缺失或重复'
    assert result['status'] == 'pending'
