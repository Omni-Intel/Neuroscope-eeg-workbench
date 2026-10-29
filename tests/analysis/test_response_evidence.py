import json

import numpy as np
import pytest
from test_assessment import fixture_record, guided_fixture

from neuroscope_eeg.analysis.assessment import assess
from neuroscope_eeg.analysis.response_evidence import ResponseEvidence, harmonic_overlaps
from neuroscope_eeg.analysis.task_assessment import classification_summary, mi_metrics
from neuroscope_eeg.reports.generate import generate_report


def test_classification_missing_evidence_is_not_rule_abstention():
    rows = [{"condition": "左手", "prediction": "左手", "analyzable": True},
            {"condition": "左手", "prediction": "不确定", "analyzable": False},
            {"condition": "右手", "prediction": "不确定", "analyzable": True},
            {"condition": "右手", "prediction": "左手", "analyzable": True}]
    summary = classification_summary(rows, ["左手", "右手"])
    assert summary["confusion"] == [[1, 0, 0, 1], [1, 0, 1, 0]]
    assert summary["accuracy"] == summary["macro_f1"] == .25
    assert summary["coverage"] == summary["selective_accuracy"] == .5
    assert summary["data_unavailable"] == summary["model_abstained"] == 1
    assert summary["precision_per_condition"]["右手"] is None
    assert np.sum(summary["confusion_row_fraction"], axis=1) == pytest.approx([1, 1])
    missing = classification_summary(rows[:1], ["左手", "右手"])
    assert missing["macro_f1"] is None and missing["balanced_accuracy"] is None
    assert missing["confusion_row_fraction"][1] == [None]*4


def test_mi_power_db_and_tfr_preserve_known_suppression():
    fs = 250
    t = np.arange(1000)/fs
    x = 10*np.sin(2*np.pi*10*t) + 6*np.sin(2*np.pi*20*t)
    baseline = np.column_stack([x, x])
    task = baseline * [1, .3]
    row = {"trial": 1, "block": 1, "condition": "左手", "clean": True, "analyzable": True,
           "completed": True, "exclusion_reason": "", **mi_metrics(baseline, task, fs)}
    assert row["mu_C4_task_power_uv2"]/row["mu_C4_baseline_power_uv2"] == pytest.approx(.09)
    assert row["mu_C4_change_db"] == pytest.approx(10*np.log10(.09))
    assert row["mu_C4_erd_pct"] == pytest.approx(-91)
    evidence = ResponseEvidence("运动想象", fs, ["C3", "C4"], [], 50)
    evidence.add_mi(row, baseline, task, task, baseline[:500])
    result = evidence.finish([row], ["左手", "右手", "静息"], 3)
    tfr = next(r for r in result["time_frequency"] if r["channel"] == "C4" and r["subset"] == "analyzable")
    idx = tfr["frequency_hz"].index(10)
    assert tfr["change_db"][idx] == pytest.approx([10*np.log10(.09)]*len(tfr["time_sec"]))
    contra = next(r for r in result["summary"] if r["metric"] == "mu_change_db" and r["channel"] == "contra-ipsi")
    assert contra["mean"] == pytest.approx(10*np.log10(.09))
    recovery = [r["value"] for r in result["trial_metrics"] if r["metric"] == "beta_recovery_erd_pct"]
    assert recovery == pytest.approx([0, 0], abs=1e-10)


def test_ssvep_overlaps_scale_invariance_and_condition_bands():
    overlaps = harmonic_overlaps([7.5, 10., 12., 15.])
    item = next(r for r in overlaps if r["target_hz"] == 7.5 and r["harmonic"] == 2)
    assert {"target_hz": 15., "harmonic": 1} in item["shared_with"]
    fs = 250
    x = np.random.default_rng(71).normal(0, 1, (1000, 2))
    x += 10*np.sin(2*np.pi*12*np.arange(1000)[:, None]/fs)
    row = {"trial": 0, "block": 1, "condition": 12., "clean": True, "completed": True, "analyzable": True, "exclusion_reason": ""}
    a, b = [ResponseEvidence("SSVEP", fs, ["O1", "O2"], [7.5, 10, 12, 15], 50) for _ in range(2)]
    a.add_ssvep(row, x)
    b.add_ssvep(row, x*2)
    for first, scaled in zip(a.metrics, b.metrics):
        if first["metric"] == "snr_db" or first["metric"].endswith("relative"):
            assert first["value"] == pytest.approx(scaled["value"])
        elif first["value"] is not None:
            assert scaled["value"] == pytest.approx(4*first["value"])
    fractions = [r["value"] for r in a.metrics if r["channel"] == "ROI" and r["metric"].endswith("relative")]
    assert sum(fractions) == pytest.approx(1)
    high = ResponseEvidence("SSVEP", fs, ["O1", "O2"], [12, 15, 24, 30], 50)
    high.add_ssvep(row, x)
    assert max(high.spectra[0]["frequency_hz"]) >= 90  # High-refresh targets retain the third harmonic.


@pytest.mark.parametrize("paradigm", ["SSVEP", "运动想象"])
def test_response_filters_and_planned_counts(tmp_path, paradigm):
    root = fixture_record(tmp_path, paradigm, repeats=2)
    data = np.fromfile(root / "raw/eeg.f32", dtype="<f4").reshape(-1, 2)
    data[800] = 1000
    data.tofile(root / "raw/eeg.f32")
    valid = np.fromfile(root / "raw/valid.u8", dtype="u1")
    valid[9*250+800] = 0
    valid.tofile(root / "raw/valid.u8")
    result = assess(root)
    task, ev = result["task"], result["task"]["response_evidence"]
    total = 8 if paradigm == "SSVEP" else 6
    all_counts = ev["retention"][0]
    assert (all_counts["observed"], all_counts["analyzable"], all_counts["clean"]) == (total, total-1, total-2)
    assert all_counts["unavailable_reasons"] and all_counts["screening_reasons"]
    assert task["all_trials"]["data_unavailable"] == 1
    assert not ev["errors"]
    first_condition = str(task["trials"][0]["condition"])
    selected = [r for r in ev["summary"] if r["condition"] == first_condition]
    assert max(r["n"] for r in selected if r["subset"] == "analyzable") == 2
    assert max(r["n"] for r in selected if r["subset"] == "screened") == 1


def test_matching_resting_windows_and_missing_baseline(tmp_path):
    _, root = guided_fixture(tmp_path / "long", "SSVEP", eyes_open_sec=20)
    report = assess(root)
    ev = report["task"]["response_evidence"]
    assert ev["resting_baseline"]["status"] == "available"
    assert ev["resting_baseline"]["subsets"]["analyzable"]["n_windows"] == 3
    assert len(ev["baseline_comparison"]) == 4*3*2  # conditions × electrodes/ROI × two subsets
    assert all(r["n_baseline_windows"] == 3 for r in ev["baseline_comparison"])
    enhanced = next(r for r in ev["baseline_comparison"] if r["channel"] == "ROI" and r["condition"] == "12.0")
    assert enhanced["change_db"] > 20
    _, short = guided_fixture(tmp_path / "short", "SSVEP")
    unavailable = assess(short)["task"]["response_evidence"]
    assert unavailable["resting_baseline"]["status"] == "unavailable"
    assert unavailable["baseline_comparison"] == []


def test_export_evidence_and_mi_figures(tmp_path):
    from docx import Document
    plan, root = guided_fixture(tmp_path, "运动想象", total_trials=9, block_trials=3, rest_sec=2)
    device = json.loads((root / "device.json").read_text(encoding="utf-8"))
    device["channel_names"] = ["c3", "c4"]
    (root / "device.json").write_text(json.dumps(device), encoding="utf-8")
    raw_before = (root / "raw/eeg.f32").read_bytes()
    output = generate_report(root)
    latest = json.loads((root / "analysis/latest.json").read_text(encoding="utf-8"))
    analysis = root / "analysis" / latest["generation"]
    report = json.loads((analysis / "assessment.json").read_text(encoding="utf-8"))
    assert report["analysis_version"] == "assessment-v1.3.0"
    assert report["configuration"]["subject_training"] is False
    assert report["task"]["response_evidence"]["errors"] == []
    assert report["task"]["counts"]["observed"] == plan.total_trials
    for name in ("trial_responses", "response_summary", "trial_retention", "time_frequency", "classification", "block_response"):
        assert (analysis / f"{name}.csv").exists()
    for name in ("response-mi-tfr", "response-mi-lateral", "response-blocks", "task-confusion"):
        assert (analysis / "figures" / f"{name}.png").exists()
    document = Document(output)
    text = "\n".join([p.text for p in document.paragraphs] + [c.text for t in document.tables for row in t.rows for c in row.cells])
    for expected in ("数据不可分析", "模型不确定", "通用信号特征", "运动想象任务响应", "对侧−同侧", "平衡准确率"):
        assert expected in text
    for expected in ("静息与正式任务的分段质量", "μ 基线／任务功率", "恢复期条件／通道", "分组响应稳定性",
                     "全程功率谱", "运动想象 ERD 时间曲线", "MI 条件功率谱", "MI 时频响应", "MI 侧化分布", "固定算法混淆矩阵"):
        assert expected in text
    assert '全程窗口质量分布' not in text
    assert len(document.inline_shapes) == 7
    ordered_text = '\n'.join(''.join(element.xpath('.//w:t/text()')) for element in document.element.body)
    anchors = ('基线与任务功率', 'μ 基线／任务功率', 'MI 条件功率谱', '功率解释',
               'μ／β 去同步及时间变化', 'μ ERD %', '运动想象 ERD 时间曲线', '时程解释',
               '响应的时频分布', 'MI 时频响应', '时频解释',
               '左右侧化', 'μ C3−C4（dB）', 'MI 侧化分布', '侧化对应', '休息早期的 β 恢复')
    positions = [ordered_text.index(anchor) for anchor in anchors]
    assert positions == sorted(positions)  # Tables, figures and interpretation stay in the same topic.
    assert '左手条件两列符号相反' in ordered_text
    assert '矩阵对应上表“全部已观察”一行' in ordered_text
    assert "ITR" not in text and "认知水平" not in text
    assert raw_before == (root / "raw/eeg.f32").read_bytes()
