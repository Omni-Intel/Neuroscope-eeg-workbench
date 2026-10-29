import json
import uuid
from datetime import UTC, datetime, timedelta
from pathlib import Path

import numpy as np
import pytest

from neuroscope_eeg.analysis.assessment import assess
from neuroscope_eeg.analysis.task_assessment import fbcca, mi_metrics
from neuroscope_eeg.io.impedance import import_measurement
from neuroscope_eeg.io.native_recording import write_json
from neuroscope_eeg.reports.generate import generate_report


def fixture_record(root, paradigm="SSVEP", repeats=2, noise=False):
    root = root / "Records" / ("recording-" + uuid.uuid4().hex)
    fs = 250
    labels = [8., 10., 12., 15.] if paradigm == "SSVEP" else ["左手", "右手", "静息"]
    total = repeats * len(labels)
    n = total*9*fs+1
    root.mkdir(parents=True)
    (root / "raw").mkdir()
    rng = np.random.default_rng(789)
    x = rng.normal(0, 2, (n, 2))
    times = np.arange(n)/fs
    if paradigm != "SSVEP":
        x += 15*np.sin(2*np.pi*10*times[:, None])+8*np.sin(2*np.pi*20*times[:, None])
    events = []
    def event(phase, trial, t, label, payload):
        events.append({"phase": phase, "monotonic_time": t+100, "wall_time": t+1000, "paradigm": paradigm,
                       "label": label, "sample_cursor": round(t*fs), "payload": {"trial": trial, "block": 1, "is_practice": False, **payload}})
    start_payload = {"ssvep_frequencies": labels if paradigm == "SSVEP" else [], "planned_formal_trials": total}
    event("start", -1, 0, paradigm, start_payload)
    for trial in range(total):
        condition = labels[trial % len(labels)]
        start = trial*9
        payload = {"target_frequency": condition, "ssvep_frequencies": labels} if paradigm == "SSVEP" else {"condition": condition}
        phases = [("cue", 0), ("flicker", 2), ("rest", 7), ("trial_complete", 9)] if paradigm == "SSVEP" else [("fixation", 0), ("cue", 2), ("imagery", 3), ("rest", 7), ("trial_complete", 9)]
        for phase, offset in phases:
            event(phase, trial, start+offset, str(condition), payload)
        if paradigm == "SSVEP" and not noise:
            a, b = (start+2)*fs, (start+7)*fs
            x[a:b] += 20*np.sin(2*np.pi*condition*times[a:b, None]) + 8*np.sin(4*np.pi*condition*times[a:b, None])
        elif paradigm != "SSVEP" and condition != "静息":
            channel = 1 if condition == "左手" else 0
            x[(start+3)*fs:(start+7)*fs, channel] *= .3
    event("stop", -1, total*9, paradigm, {"completed_formal_trials": total, "missed_frames_estimate": 0})
    now = datetime(2026, 10, 28, 0, 0, tzinfo=UTC)
    run = {"session_id": root.name.removeprefix("recording-"), "participant_id": "00000001", "paradigm": paradigm,
           "paradigm_code": "TEST-20261028", "unique_name": "ASSESS-20261028", "task": "ssvep" if paradigm == "SSVEP" else "motorimagery",
           "started_at": now.isoformat(), "ended_at": (now+timedelta(seconds=total*9)).isoformat(),
           "external_recording": False, "simulated": True, "recording_status": "complete", "stop_reason": "completed",
           "samples": n, "invalid_samples": 0, "continuity_errors": [], "protocol_start_payload": start_payload,
           "acquisition_software": {"name": "NeuroScope", "version": "test"}}
    device = {"source_id": "fixture", "source_type": "simulated", "sfreq": fs, "channel_names": ["O1", "O2"] if paradigm == "SSVEP" else ["C3", "C4"],
              "channel_types": ["eeg"]*2, "channel_units": ["uV"]*2, "extra": {}}
    for name, value in (("run", run), ("device", device), ("config", {"profile": {"confirmed": True, "eeg_reference": "synthetic", "power_line_frequency": 50}})):
        write_json(root / (name+".json"), value)
    x.astype("<f4").tofile(root / "raw/eeg.f32")
    times.astype("<f8").tofile(root / "raw/timestamps.f64")
    np.arange(n, dtype="<i8").tofile(root / "raw/sample_indices.i64")
    np.ones(n, dtype="u1").tofile(root / "raw/valid.u8")
    (root / "raw/chunks.jsonl").write_text(json.dumps({"samples": n})+"\n", encoding="utf-8")
    (root / "events.jsonl").write_text("\n".join(json.dumps(e, ensure_ascii=False) for e in events)+"\n", encoding="utf-8")
    (root / "behavior.csv").write_text("phase,response\n", encoding="utf-8")
    return root


def test_ssvep_known_targets_noise_and_reproducibility(tmp_path):
    root = fixture_record(tmp_path / "positive")
    a, b = assess(root), assess(root)
    assert a == b
    assert a["task"]["all_trials"]["accuracy"] == 1
    assert a["task"]["counts"] == {"observed": 8, "completed": 8, "analyzable": 8, "clean": 8}
    assert all(r["snr_db"] > 20 for r in a["task"]["trials"])
    noise = assess(fixture_record(tmp_path / "noise", repeats=10, noise=True))
    assert noise["task"]["all_trials"]["accuracy"] < .5
    # Prediction itself accepts no target label. Changing labels cannot train it.
    x = np.sin(2*np.pi*12*np.arange(1000)/250)[:, None]
    assert fbcca(x, 250, [8, 10, 12, 15])[0] == 2


def test_mi_erd_sign_and_side_rule(tmp_path):
    report = assess(fixture_record(tmp_path / "mi", "运动想象"))
    assert report["task"]["all_trials"]["accuracy"] == 1
    erd = report["task"]["erd"]
    assert erd["左手"]["mu_C4_erd_pct"]["mean"] < -80
    assert erd["右手"]["mu_C3_erd_pct"]["mean"] < -80
    assert abs(erd["静息"]["mu_C3_erd_pct"]["mean"]) < 10
    assert report["task"]["timecourses"]
    with pytest.raises(ValueError, match="基线"):
        mi_metrics(np.zeros((400, 2)), np.ones((750, 2)), 250)


def guided_fixture(root, paradigm, **overrides):
    from dataclasses import replace
    from neuroscope_eeg.protocol_config import AssessmentProtocol
    from neuroscope_eeg.desktop.guided_protocol import protocol_steps

    root = fixture_record(root, paradigm, repeats=1)
    base = AssessmentProtocol.defaults(paradigm)
    if paradigm == 'SSVEP':
        base = replace(base, ssvep_mode='single_target_v2')
    plan = replace(base, total_trials=base.conditions*2, block_trials=base.conditions,
                   task_sec=6, cue_sec=1.5, fixation_sec=2.5, rest_sec=1,
                   eyes_open_sec=5, eyes_closed_sec=5, close_prepare_sec=3, open_recovery_sec=3, block_rest_sec=5)
    plan = replace(plan, **overrides).validate()
    fs, elapsed, events, intervals = 250, 0., [], []
    frequencies = list(map(float, range(8,16))) if plan.eight_target else [8., 10., 12., 15.]
    def add(phase, payload, t):
        events.append({"phase": phase, "paradigm": paradigm, "label": "", "sample_cursor": round(t*fs),
                       "monotonic_time": t+100, "wall_time": t+1000, "payload": payload})
    add("start", {"assessment_protocol": plan.to_dict()}, 0)
    for i, step in enumerate(protocol_steps(plan)):
        payload = {"guided_step": i, "planned_duration_s": step.duration,
                   "segment": "baseline" if step.phase.startswith("baseline_") else "instruction", **step.payload}
        if "trial" in payload:
            if paradigm == "SSVEP":
                payload.update(target_frequency=frequencies[payload["condition_index"]], ssvep_frequencies=frequencies)
            else:
                payload["condition"] = ("左手", "右手", "静息")[payload["condition_index"]]
        add(step.phase, payload, elapsed)
        end = elapsed + (step.duration if step.duration is not None else 1)
        intervals.append((step.phase, payload, elapsed, end))
        if step.phase == "rest":
            add("trial_complete", payload, end)
        if step.phase in {"baseline_eyes_open", "baseline_eyes_closed"}:
            add(step.phase + "_end", {"segment": "baseline"}, end)
        elapsed = end
    add("stop", {"completed_formal_trials": plan.total_trials, "missed_frames_estimate": 0}, elapsed)
    n = round(elapsed*fs)+1
    times = np.arange(n)/fs
    data = np.random.default_rng(321).normal(0, 2, (n, 2)) + 10*np.sin(2*np.pi*10*times[:, None])
    for phase, payload, start, end in intervals:
        a, b = round(start*fs), round(end*fs)
        if phase == "flicker":
            data[a:b] += 30*np.sin(2*np.pi*payload["target_frequency"]*times[a:b, None])
        elif phase == "imagery" and payload["condition"] != "静息":
            data[a:b, 1 if payload["condition"] == "左手" else 0] *= .3
    run = json.loads((root / "run.json").read_text(encoding="utf-8"))
    run.update(samples=n, protocol_start_payload={"assessment_protocol": plan.to_dict(), "planned_formal_trials": plan.total_trials,
                                                "ssvep_frequencies": frequencies})
    write_json(root / "run.json", run)
    config = json.loads((root / "config.json").read_text(encoding="utf-8"))
    write_json(root / "config.json", {**config, "assessment_protocol": plan.to_dict()})
    data.astype("<f4").tofile(root / "raw/eeg.f32")
    times.astype("<f8").tofile(root / "raw/timestamps.f64")
    np.arange(n, dtype="<i8").tofile(root / "raw/sample_indices.i64")
    np.ones(n, dtype="u1").tofile(root / "raw/valid.u8")
    (root / "events.jsonl").write_text("\n".join(json.dumps(e, ensure_ascii=False) for e in events)+"\n", encoding="utf-8")
    (root / "raw/chunks.jsonl").write_text(json.dumps({"samples": n})+"\n", encoding="utf-8")
    return plan, root


@pytest.mark.parametrize("paradigm", ["SSVEP", "运动想象"])
def test_custom_timing_and_baseline_excluded_from_task(tmp_path, paradigm):
    root = tmp_path / "guided"
    plan, root = guided_fixture(root, paradigm)
    result = assess(root)
    assert result["task"]["counts"]["analyzable"] == plan.total_trials
    assert result["task"]["counts"]["observed"] == plan.total_trials
    assert [r["label"] for r in result["segment_quality"]] == ["睁眼静息", "闭眼静息", "正式第 1 组", "正式第 2 组"]
    assert result["segment_quality"][0]["duration_sec"] == 5
    assert result["segment_quality"][2]["duration_sec"] == sum(t for _, t in plan.phases)*plan.block_trials
    assert result["configuration"]["mi_task_relative_imagery_sec"] == [.5, 5.5]
    assert result["configuration"]["mi_baseline_relative_fixation_sec"] == [.2, 2.3]
    from docx import Document
    from neuroscope_eeg.reports.word_report import write_word
    write_word(result, tmp_path / "custom.docx", [])
    document = Document(tmp_path / "custom.docx")
    text = "\n".join(p.text for p in document.paragraphs)
    assert "0.5–5.5" in text and "0.5–4.5" not in text
    assert "实验完成与数据可用性" in text and "phases.csv" in text
    # Tampered saved configuration is never silently scored with old defaults.
    config = json.loads((root / "config.json").read_text(encoding="utf-8"))
    config["assessment_protocol"]["task_sec"] = 5
    write_json(root / "config.json", config)
    result = assess(root)
    assert result["task"]["status"] == "unavailable" and "不一致" in result["task"]["reason"]


def test_ssvep_requires_positioning_event_only_for_new_protocol(tmp_path):
    plan, root = guided_fixture(tmp_path / 'new', 'SSVEP', total_trials=4, block_trials=4)
    events = [json.loads(line) for line in (root / 'events.jsonl').read_text(encoding='utf-8').splitlines()]
    first_trial = plan.practice_per_condition * plan.conditions
    events = [e for e in events if not (e['phase'] == 'prep' and e['payload'].get('trial') == first_trial)]
    (root / 'events.jsonl').write_text('\n'.join(json.dumps(e) for e in events)+'\n', encoding='utf-8')
    result = assess(root)
    assert result['task']['counts']['analyzable'] == 3
    assert '阶段记录' in result['task']['trials'][0]['exclusion_reason']
    legacy, root = guided_fixture(tmp_path / 'legacy', 'SSVEP', total_trials=4, block_trials=4,
                                  ssvep_mode='simultaneous_quadrants_v1')
    result = assess(root)
    assert result['task']['counts']['analyzable'] == 4
    assert result['protocol']['settings']['version'] == 'guided-assessment-v1'


@pytest.mark.parametrize("damage", ["invalid", "gap", "truncated_trial", "unit", "electrode", "negative_cursor"])
def test_no_fabricated_task_scores(tmp_path, damage):
    root = fixture_record(tmp_path / damage, repeats=1)
    if damage == "invalid":
        values = np.fromfile(root / "raw/valid.u8", dtype="u1")
        values[800] = 0
        values.tofile(root / "raw/valid.u8")
    elif damage == "gap":
        values = np.fromfile(root / "raw/sample_indices.i64", dtype="<i8")
        values[800:] += 20
        values.tofile(root / "raw/sample_indices.i64")
    elif damage in {"truncated_trial", "negative_cursor"}:
        events = [json.loads(line) for line in (root / "events.jsonl").read_text(encoding="utf-8").splitlines()]
        if damage == "truncated_trial":
            events = [e for e in events if not (e["phase"] == "trial_complete" and e["payload"]["trial"] == 0)]
        else:
            events[1]["sample_cursor"] = -1
        (root / "events.jsonl").write_text("\n".join(json.dumps(e) for e in events), encoding="utf-8")
    else:
        device = json.loads((root / "device.json").read_text(encoding="utf-8"))
        device["channel_units" if damage == "unit" else "channel_names"] = ["counts", "counts"] if damage == "unit" else ["F3", "F4"]
        write_json(root / "device.json", device)
    report = assess(root)
    if damage in {"invalid", "gap", "truncated_trial"}:
        assert report["task"]["counts"]["analyzable"] == 3
        assert report["task"]["all_trials"]["n"] == 4
        assert report["task"]["all_trials"]["accuracy"] == .75
        assert report["task"]["trials"][0]["exclusion_reason"]
    else:
        assert report["task"]["status"] == "unavailable"
        assert report["channels"]


def test_full_record_qc_not_last_minute_and_line_before_notch(tmp_path):
    root = fixture_record(tmp_path / "qc", repeats=3)
    data = np.memmap(root / "raw/eeg.f32", dtype="<f4", mode="r+", shape=(108*250+1, 2))
    data[:1000, 0] = 0
    data[1000:2000, 0] += 400*np.sin(2*np.pi*50*np.arange(1000)/250)
    data.flush()
    del data
    report = assess(root)
    windows = [r for r in report["windows"] if r["channel"] == "O1"]
    assert windows[0]["flat_suspected"]
    assert windows[1]["line_fraction"] > .9
    assert report["channels"][0]["clipping_fraction_median"] is None
    assert report["configuration"]["notch_applied"] is False


def test_impedance_identity_and_no_unit_conversion(tmp_path):
    root = fixture_record(tmp_path / "run", repeats=1)
    source = tmp_path / "impedance.json"
    value = {"source_id": "fixture", "measured_at": "2026-10-28T07:59:00+08:00", "stage": "pre", "unit": "CQ",
             "method": "manual source", "channels": [{"channel": "O1", "value": 3}]}
    write_json(source, value)
    result = import_measurement(root, source)
    saved = json.loads(result.read_text(encoding="utf-8"))
    assert saved["unit"] == "CQ" and saved["channels"][0]["value"] == 3
    assert len(assess(root)["impedance"]) == 1
    value["source_id"] = "other"
    write_json(source, value)
    with pytest.raises(ValueError, match="设备"):
        import_measurement(root, source)


def test_word_outputs_versions_and_failure_retains_raw(tmp_path):
    root = fixture_record(tmp_path / "报告中文", repeats=1)
    before = (root / "raw/eeg.f32").read_bytes()
    first = generate_report(root)
    second = generate_report(root)
    assert first != second and first.exists() and second.exists()
    from docx import Document
    doc = Document(first)
    assert first.stem == doc.core_properties.title == doc.paragraphs[0].text
    assert first.name == '00000001 SSVEP脑电信号质量评估报告（模拟数据源）.docx'
    captions = '\n'.join(p.text for p in doc.paragraphs)
    for expected in ('全程功率谱', 'SSVEP 条件频谱与 SNR 谱', 'SSVEP 逐通道响应分布', '固定算法混淆矩阵'):
        assert expected in captions
    assert '全程窗口质量分布' not in captions
    assert len(doc.inline_shapes) == 4  # No screening image; numeric window evidence remains.
    ordered_text = '\n'.join(''.join(element.xpath('.//w:t/text()')) for element in doc.element.body)
    anchors = ('频率响应与谐波', '基频 SNR', 'SSVEP 条件频谱与 SNR 谱', '图表对应',
               '与独立睁眼静息的对照', '逐通道响应分布', 'SSVEP 逐通道响应分布', '通道解释')
    positions = [ordered_text.index(anchor) for anchor in anchors]
    assert positions == sorted(positions)
    assert '图上的峰值不要求等于表格数值' in ordered_text
    latest = json.loads((root / 'analysis/latest.json').read_text(encoding='utf-8'))
    assert Path(latest['report']) == second
    for file in ("assessment.json", "channels.csv", "trials.csv", "windows.csv", "timecourses.csv", "status.json"):
        assert (root / "analysis" / first.parent.name / file).exists()
    assert not list(root.rglob("*.html"))
    assert before == (root / "raw/eeg.f32").read_bytes()
    (root / "raw/valid.u8").write_bytes(b"bad")
    with pytest.raises(ValueError, match="长度"):
        generate_report(root)
    assert first.exists()
    assert list(root.glob("analysis/.pending-*/error.log"))


@pytest.mark.parametrize("paradigm", ["SSVEP", "运动想象"])
def test_reference_report_structure_and_no_borrowed_results(tmp_path, paradigm):
    from docx import Document

    from neuroscope_eeg.reports.word_report import write_word
    report = assess(fixture_record(tmp_path / "source", paradigm, repeats=1))
    report["report_generated_at"] = "2026-10-28T23:30:00+00:00"
    output = tmp_path / "report.docx"
    write_word(report, output, [])
    doc = Document(output)
    assert [row.cells[0].text for row in doc.tables[0].rows] == [
        "项目", "被试", "有效任务", "采集系统", "分析电极", "重参考", "预处理", "事件同步", "分析日期"]
    assert "2026-10-29" in doc.tables[0].cell(8, 1).text
    assert doc.paragraphs[0].style.name == "Title"
    assert doc.styles["Title"].font.size.pt == 27
    assert doc.styles["Normal"].font.name == "Microsoft YaHei"
    assert doc.sections[0].left_margin.inches == 1
    text = "\n".join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for row in t.rows for c in row.cells])
    for expected in ["实验完成与数据可用性", "结果摘要", "最终综合结论", "仅验证流程"]:
        assert expected in text
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert headings.index("实验完成与数据可用性") < headings.index("信号质量与电极接触")
    assert headings.index("信号质量与电极接触") < headings.index("阻抗与接触质量") < headings.index(paradigm + "任务响应")
    assert text.count("结果摘要") == 1
    for absent in ["本次判定", "总体判断", "核心结论", "MI 使用", "MI 时频", "采集开始", "开始采集时间", "Fpz 重参考", "87.5%", "0.813", "hw 的", "complete / completed"]:
        assert absent not in text


def test_unavailable_mi_report_explains_exclusions_without_zero_accuracy(tmp_path):
    from docx import Document

    from neuroscope_eeg.reports.word_report import write_word
    root = fixture_record(tmp_path / "source", "运动想象", repeats=1)
    events = [json.loads(line) for line in (root / "events.jsonl").read_text(encoding="utf-8").splitlines()]
    for event in events:
        event["sample_cursor"] = round(event["sample_cursor"] * .95)
    (root / "events.jsonl").write_text("\n".join(json.dumps(e, ensure_ascii=False) for e in events), encoding="utf-8")
    report = assess(root)
    output = tmp_path / "unavailable.docx"
    write_word(report, output, [])
    doc = Document(output)
    text = "\n".join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for row in t.rows for c in row.cells])
    counts_table = next(t for t in doc.tables if [c.text for c in t.rows[0].cells] == ["条件", "计划", "完成", "可分析／观察", "筛查后／可分析"])
    totals = next(row for row in counts_table.rows if row.cells[0].text == "全部")
    assert [totals.cells[i].text for i in (1, 2, 3)] == ["3", "3", "0/3"]
    assert "实际阶段时长偏离计划" in text
    assert "不可分析不应解读为设备表现差" in text
    assert "准确率为 0.0%" not in text  # Unavailable task scores must not be presented as failures.


def test_external_edf_waits_for_alignment_and_checks_hash(tmp_path):
    import pyedflib

    from neuroscope_eeg.io.external_recording import associate_external
    root = fixture_record(tmp_path / "external", repeats=1)
    run = json.loads((root / "run.json").read_text(encoding="utf-8"))
    run["external_recording"] = True
    write_json(root / "run.json", run)
    waiting = assess(root)
    assert waiting["task"]["status"] == "unavailable"
    source = tmp_path / "emotiv.edf"
    data = np.fromfile(root / "raw/eeg.f32", dtype="<f4").reshape(-1, 2)[:-1]
    with pyedflib.EdfWriter(str(source), 2) as writer:
        writer.setSignalHeaders([{"label": name, "dimension": "uV", "sample_frequency": 250,
                                  "physical_min": -100, "physical_max": 100, "digital_min": -32768, "digital_max": 32767}
                                 for name in ("O1", "O2")])
        writer.writeSamples([data[:, i] for i in range(2)])
    associate_external(root, source, operator_confirmed=True)
    write_json(root / "analysis_profile.json", {"confirmed": True, "channel_types": ["eeg", "eeg"]})
    before = assess(root)
    assert before["channels"] and before["task"]["status"] == "unavailable"
    assert "对齐" in before["task"]["reason"]
    events = [json.loads(line) for line in (root / "events.jsonl").read_text(encoding="utf-8").splitlines()]
    write_json(root / "alignment.json", {"session_id": run["session_id"], "method": "fixture_known_time",
               "evidence": "Synthetic fixture has exactly matched sample times",
               "events": [{"event_index": i, "onset": e["sample_cursor"]/250, "duration": 0} for i, e in enumerate(events)]})
    after = assess(root)
    assert after["task"]["all_trials"]["accuracy"] == 1
    with (root / "external/recording.edf").open("ab") as f:
        f.write(b"tampered")
    with pytest.raises(ValueError, match="校验"):
        assess(root)


def test_noisy_trial_not_silently_removed_from_accuracy(tmp_path):
    root = fixture_record(tmp_path / "artifact", repeats=1)
    data = np.memmap(root / "raw/eeg.f32", dtype="<f4", mode="r+", shape=(36*250+1, 2))
    data[800] = 1000
    data.flush()
    del data
    task = assess(root)["task"]
    assert task["all_trials"]["n"] == 4
    assert task["analyzable_trials"]["n"] == 4
    assert task["clean_trials"]["n"] == 3
    assert "大幅值" in task["trials"][0]["exclusion_reason"]
