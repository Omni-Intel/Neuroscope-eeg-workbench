"""Reusable styles and components distilled from the user's August 21 report."""
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

TEMPLATE = Path(__file__).with_name("templates") / "eeg-reference.docx"
LAYOUT_VERSION = "reference-8.21-v3"


class ReportLayout:
    def __init__(self):
        self.doc = Document(TEMPLATE)
        self.table_example = deepcopy(self.doc.tables[0]._tbl)
        self.box_example = deepcopy(self.doc.tables[1]._tbl)
        for table in list(self.doc.tables):
            table._tbl.getparent().remove(table._tbl)
        normal = self.doc.styles["Normal"].paragraph_format
        normal.line_spacing = 1.1
        normal.space_after = Pt(6)
        self.figure_number = 0

    def paragraph(self, text, lead=None):
        p = self.doc.add_paragraph()
        if lead:
            p.add_run(lead + "：").bold = True
        p.add_run(text)
        return p

    def table(self, headers, rows, widths=None, *, compact=False):
        t = self.doc.add_table(rows=1, cols=len(headers))
        t.autofit = False
        raw = widths or [1] * len(headers)
        widths = [6.5 * w / sum(raw) for w in raw]
        t._tbl.remove(t._tbl.tblPr)
        t._tbl.insert(0, deepcopy(self.table_example.tblPr))
        t._tbl.tblPr.append(OxmlElement("w:tblLayout"))
        t._tbl.tblPr[-1].set(qn("w:type"), "fixed")
        for col, width in zip(t.columns, widths):
            col.width = Inches(width)
        for _ in rows:
            t.add_row()
        for index, (row, values) in enumerate(zip(t.rows, [headers, *rows])):
            prototype = self.table_example.tr_lst[0 if index == 0 else 1].tc_lst[0]
            for cell, text, width in zip(row.cells, values, widths):
                cell._tc.remove(cell._tc.tcPr)
                props = deepcopy(prototype.tcPr)
                for old in list(props.findall(qn("w:tcW"))):
                    props.remove(old)
                cell._tc.insert(0, props)
                cell.width = Inches(width)
                p = cell.paragraphs[0]
                p._p.insert(0, deepcopy(prototype.p_lst[0].pPr))
                # Keep small tables intact and avoid a lone first/last body row on a page.
                p.paragraph_format.keep_with_next = (index < len(rows) if len(rows) <= 4
                                                    else index in {0, 1, len(rows)-1})
                p.paragraph_format.keep_together = True
                if compact:
                    grid = OxmlElement("w:snapToGrid")
                    grid.set(qn("w:val"), "0")
                    p._p.get_or_add_pPr().append(grid)
                r = p.add_run(str(text))
                source_run = next(r for r in prototype.p_lst[0].r_lst if r.rPr is not None)
                r._r.insert(0, deepcopy(source_run.rPr))
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        t.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.line_spacing = Pt(2)
        p.paragraph_format.widow_control = False
        return t

    def box(self, lead, text):
        element = deepcopy(self.box_example)
        self.doc.element.body.insert_element_before(element, "w:sectPr")
        cell = self.doc.tables[-1].cell(0, 0)
        p = cell.paragraphs[0]
        source = next(r for r in p.runs if r._r.rPr is not None)
        props = deepcopy(source._r.rPr)
        p.clear()
        for content, bold in [(lead + "：", True), (text, False)]:
            run = p.add_run(content)
            run._r.insert(0, deepcopy(props))
            run.bold = bold
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.space_before = Pt(5)
        cell._tc.getparent().get_or_add_trPr().append(OxmlElement("w:cantSplit"))

    def figure(self, title, path, *, width=6.5):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_after = Pt(2)
        p.add_run().add_picture(str(path), width=Inches(width))
        self.figure_number += 1
        caption = self.doc.add_paragraph(f"图 {self.figure_number}　{title}")
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.space_after = Pt(8)
        for r in caption.runs:
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor.from_string("666666")
