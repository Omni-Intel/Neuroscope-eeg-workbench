"""Evidence-led task report with selected tables and complete diagnostic figures."""
from collections import Counter
from datetime import datetime, timedelta, timezone

import numpy as np
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from neuroscope_eeg.reports.layout import ReportLayout


def write_compact(report, destination, figures, title, device_label, shown):
    layout = ReportLayout()
    doc, p, table = layout.doc, layout.paragraph, layout.table
    def table(headers, rows, widths=None):
        return layout.table(headers, rows, widths, compact=True)
    doc.styles['Title'].paragraph_format.space_before = Pt(0)
    doc.styles['Subtitle'].paragraph_format.space_after = Pt(8)
    doc.styles['Normal'].paragraph_format.space_after = Pt(4)
    doc.styles['Normal'].paragraph_format.line_spacing = 1.05
    # Use the declared line spacing instead of rounding Chinese text up to the template's 18 pt grid.
    snap = OxmlElement('w:snapToGrid')
    snap.set(qn('w:val'), '0')
    doc.styles['Normal']._element.get_or_add_pPr().append(snap)
    run, dev, task, protocol = (report[k] for k in ('run', 'device', 'task', 'protocol'))
    ssvep = run['paradigm'] == 'SSVEP'
    simulated = run.get('simulated', False)
    evidence = task.get('response_evidence', {})
    profile = report['profile']
    paired_groups = []
    def group_start():
        return len(doc.element.body) - 1

    def group_end(start):
        paired_groups.append(list(doc.element.body)[start:-1])

    def figure_named(name):
        for label, path in figures:
            if path.name == name:
                layout.figure(label, path)

    doc.core_properties.title = title
    doc.add_paragraph(title, 'Title')
    doc.add_paragraph('信号质量、通用特征与任务响应评估', 'Subtitle')
    generated = report.get('report_generated_at')
    date = datetime.fromisoformat(generated).astimezone(timezone(timedelta(hours=8))).date().isoformat() if generated else '未记录'
    table(['项目', '内容'], [
        ('被试', run['participant_id'] + ('（模拟流程测试编号）' if simulated else '')),
        ('有效任务', run['paradigm']),
        ('采集系统', f"{device_label}；{len(dev['channel_names'])} 通道；{dev['sfreq']:g} Hz；单位 {'、'.join(dict.fromkeys(dev['channel_units']))}"),
        ('分析电极', '、'.join(task.get('selected_channels', [])) or '未确认专项分析电极'),
        ('重参考', '保留采集参考；' + (profile.get('eeg_reference') or '参考方式未确认')),
        ('预处理', '通用质量：原始数据 4 秒窗，不陷波；通用特征另用 1–40 Hz、250 Hz；任务按保存协议分析。'),
        ('事件同步', '、'.join({'software_cursor_unverified':'样本索引近似对齐，未校准', 'external_alignment':'外部对齐依据'}.get(m,m) for m in task.get('timing_methods', [])) + '；软件对齐不等于已验证物理刺激时刻'),
        ('分析日期', date + '（北京时间）'),
    ], [1.05, 5.45])
    p(('本次为模拟 / 回放数据，仅验证流程，不代表真实设备性能。' if simulated else '本报告描述本次记录，不能据此确定设备总体等级或被试能力。')
      + ('已获得可计算的任务响应和固定算法结果。' if task['status'] == 'available' else '任务指标暂不可用：' + task.get('reason', '证据不足')), '结果摘要')
    if dev['source_type'] == 'simulated' and dev['source_id'] == 'simulated':
        p('当前模拟源独立循环生成信号，不随提示产生对应 SSVEP 或左右手运动想象响应；准确率与响应数值仅用于检查软件流程。')
    doc.add_heading('实验完成与数据可用性', 1)
    rows = evidence.get('retention', [])
    if rows:
        table(['条件', '计划', '完成', '可分析／观察', '筛查后／可分析'], [
            (r['condition'], shown(r['planned']), r['completed'], f"{r['analyzable']}/{r['observed']}", f"{r['clean']}/{r['analyzable']}") for r in rows], [1.1,.8,.8,1.9,1.9])
    else:
        table(['计划正式试次', '实际完成', '可分析'], [(shown(protocol['planned']), shown(protocol['completed']), shown(task.get('counts', {}).get('analyzable')))])
    reasons = Counter(r.get('exclusion_reason') or '原因未记录' for r in task.get('trials', []) if not r.get('analyzable'))
    if reasons:
        p('；'.join(f'{reason}：{n} 次' for reason, n in reasons.items()), '不可分析原因')
    screening = rows[0].get('screening_reasons', {}) if rows else {}
    if screening:
        p('；'.join(f'{reason}：{n} 次' for reason, n in screening.items()) + '。同一试次可能命中多条规则。', '幅值筛查未通过原因')
    p('练习不计分；完成状态来自程序事件，不证明实际注视或想象。筛查后仅表示通过现有幅值规则。')

    from neuroscope_eeg.reports.protocol_section import write_protocol_section, write_segment_quality_section
    write_protocol_section(layout, report, shown)
    p("阶段事件与完整时间明细见 phases.csv；实测时间不等于已验证的物理刺激时刻。")

    doc.add_heading('信号质量与电极接触', 1)
    integrity = report['integrity']
    channels = [c for c in report['channels'] if c['type'] == 'eeg']
    def range_value(key, percent=False):
        vals = [c[key] for c in channels if c.get(key) is not None]
        return f'{shown(min(vals),percent)}～{shown(max(vals),percent)}' if vals else '不可用'
    table(['检查项', '本次结果'], [
        ('记录时长／有效样本', f"{shown(integrity['duration_sec'])} 秒／{shown(integrity['samples']-integrity['invalid_samples']) if integrity.get('invalid_samples') is not None else '不可用'}"),
        ('样本序号间断／长时间戳间隔', f"{shown(integrity.get('sample_index_gaps'))}／{shown(integrity.get('timestamp_long_intervals'))}"),
        ('硬件发送错误', shown(protocol.get('hardware_send_errors'))),
        ('非有限样本通道数', sum(c['finite_fraction'] < 1 for c in channels) if channels else '无已确认 EEG 通道'),
        ('RMS／峰峰值通道范围（μV）', range_value('rms_ac_median') + '／' + range_value('peak_to_peak_median') if channels and all(c['unit'] in {'uV','µV','μV'} for c in channels) else '单位不统一，见 channels.csv'),
        ('平直窗／大幅值窗比例通道范围', range_value('flat_fraction',True) + '／' + range_value('large_amplitude_fraction',True)),
        ('工频比例通道范围', range_value('line_fraction_median',True)),
    ], [3.2,3.3])
    p('幅值为通道内 4 秒窗统计的中位数；平直 <0.1 μV、大幅值 >300 μV 是工程候选阈值。范围不表示整段污染比例；时间间隔不等于已确认丢包，工频未确认时不计算。')
    from neuroscope_eeg.reports.timing_section import write_timing_section
    write_timing_section(layout, report, shown)
    flagged_quality = [c for c in channels if (c.get('flat_fraction') or 0)>0 or (c.get('large_amplitude_fraction') or 0)>0 or c['finite_fraction']<1 or (c.get('clipping_fraction_median') or 0)>0]
    if flagged_quality:
        table(['异常线索通道', '非有限样本', '平直窗', '大幅值窗', '量程触及'], [
            (c['channel'], shown(1-c['finite_fraction'],True), shown(c.get('flat_fraction'),True),
             shown(c.get('large_amplitude_fraction'),True), shown(c.get('clipping_fraction_median'),True))
            for c in flagged_quality], [1.5,1.25,1.25,1.25,1.25])
        p('以上为筛查线索，不能确认眼电、肌电等具体来源；量程未确认时不估计触及比例。')
    for label,path in figures:
        if label == '全程功率谱':
            layout.figure(label,path,width=5)
            p('图为各通道有效窗口的平均功率谱，上表 RMS／峰峰值则是逐窗幅值的中位数，两者不是同一个统计量。此图包含记录中的练习与等待，仅用于观察全程频谱和干扰。', '频谱解释')
    write_segment_quality_section(layout, report, shown)
    doc.add_heading('阻抗与接触质量', 2)
    if report['impedance']:
        for measurement_index, item in enumerate(report['impedance'], 1):
            vals = [r.get('value') for r in item.get('channels', []) if isinstance(r.get('value'), (int,float))]
            p(f"{item.get('method','外部导入')}；{item.get('measured_at','时间未提供')}；{item.get('stage','阶段未提供')}；"
              f"{len(item.get('channels',[]))} 通道；范围 {shown(min(vals)) if vals else '不可用'}～{shown(max(vals)) if vals else '不可用'} {item.get('unit','单位未确认')}。仅代表测量时点。")
            table(['通道', '测量值', '单位', '厂商或导入标记'], [(r['channel'], shown(r.get('value')), item.get('unit','未确认'), r.get('status','')) for r in item.get('channels',[])], [1.2,1.2,1.2,2.9])
            figure_named(f'impedance-{measurement_index}.png')
            p('图与上表使用同一次测量；无数值的通道只保留在表中。不同单位或厂商等级不合并比较，不设跨品牌合格线。', '接触解释')
    else:
        p('未获得本次实际阻抗或接触测量，暂不能评价电极接触。')
    if dev['source_type'] == 'emotiv_external':
        p('EMOTIV CQ／EQ 等级不等同于 kΩ，不作阻抗单位换算。')
    doc.add_heading('通用信号特征', 2)
    general = report.get('general_features', {})
    if general.get('status') == 'available':
        keys = ('spect_entropy','samp_entropy','hjorth_mobility','spect_slope')
        segments = list(dict.fromkeys(r['segment'] for r in general['summary']))
        table(['区段', '谱熵', '样本熵', 'Hjorth mobility', '谱斜率'], [
            (seg, *(shown(next(r['value'] for r in general['summary'] if r['segment']==seg and r['metric']==k)) for k in keys)) for seg in segments], [1.1,1,1,1.8,1.6])
        p('入选通道窗数：' + '；'.join(seg + ' ' + str(sum(r['selected'] for r in general['selection'] if r['segment']==seg)) for seg in segments) + '。按通道分别计数，非独立被试样本量。')
        p('分析电极的独立 4 秒窗经 1–40 Hz 滤波并重采样至 250 Hz，每区段每通道最多均匀选取 30 个合格窗；先取通道内中位数，再取跨通道中位数。正式区段含提示、任务及试次内休息；排除原因见 general_selection.csv。')
        p('熵描述规则性，Hjorth 描述变化特征；谱斜率为含振荡峰的 2–40 Hz 拟合，不是分离后的非周期指数。均无“越高越好”的统一方向，不据此判定脑电有效或设备合格。扩展复杂度指标仅存明细。')
    else:
        p(general.get('reason','缺少可分析数据'), '通用特征暂不可用')

    doc.add_heading('方法与数据追溯', 2)
    p(f"记录标识 {run['session_id']}；分析版本 {report['analysis_version']}。")
    p('通用特征参考 N-back 工具包的 mne-features／antropy 指标体系；滤波带宽、参考方式和分段对象不同，不直接与原 N-back 数值比较。实际参数与库版本见 assessment.json。')
    p('同次分析目录提供 CSV：channels/windows（质量）、trials/trial_retention（试次与排除）、general_*（通用特征）、response_summary/trial_responses（响应）、classification（分类）；完整图像见 figures。')
    doc.add_heading(run['paradigm'] + '任务响应', 1).paragraph_format.page_break_before = True
    summary = evidence.get('summary', [])
    def metric(cond, channel, key, subset='analyzable', harmonic=None):
        return next((r for r in summary if r['condition']==str(cond) and r['channel']==channel and r['metric']==key and r['subset']==subset and r['harmonic']==harmonic), {})
    def estimate(row):
        if row.get('mean') is None:
            return '不可用'
        ci=row.get('ci95')
        return shown(row['mean']) + (f" [{shown(ci[0])}, {shown(ci[1])}]" if ci else ' [区间不足]')
    conditions = [r['condition'] for r in rows if r['condition'] != '全部']
    if not summary:
        p(task.get('reason') or '没有可用的补充响应结果。')
        p('不可分析不应解读为设备表现差。')
    elif ssvep:
        doc.add_heading('频率响应与谐波', 2)
        lo,hi=report['configuration']['ssvep_window_relative_flicker_sec']
        p(f'取闪烁后 {lo:g}–{hi:g} 秒；SNR 为目标频带与邻近背景功率比的 10log10，排除其他目标、谐波及已知工频邻域。')
        if len(run.get('protocol_start_payload', {}).get('ssvep_frequencies', [])) == 8:
            p('八频率间隔 1 Hz：背景取中心 ±2 Hz 内、距中心至少 0.5 Hz 的频点；排除各目标及谐波 ±max(0.25 Hz, 频谱分辨率)。至少保留 3 个背景频点，否则不计算 SNR。此口径与旧四分类不同，不能直接比较 SNR 数值。')
        table(['Hz／范围', '基频 n', '基频 SNR [95%区间]', '二次 SNR', '三次 SNR'], [
            (f'{c}／{label}',metric(c,'ROI','snr_db',subset,1).get('n',0),estimate(metric(c,'ROI','snr_db',subset,1)),
             shown(metric(c,'ROI','snr_db',subset,2).get('mean')),shown(metric(c,'ROI','snr_db',subset,3).get('mean')))
            for c in conditions for subset,label in [('analyzable','可分析'),('screened','筛查后')]], [1.3,.6,2.2,1.2,1.2])
        p('SNR 单位 dB；基频与谐波的有效样本数可能不同，二／三次谐波 n 与区间见 response_summary.csv；不足背景频点时显示不可用。')
        figure_named('response-ssvep-spectrum.png')
        p('表为逐试次目标频带 SNR（dB）的均值；图为可分析试次平均功率谱，并逐频点计算 SNR。目标频带与单频点、先取比值与先平均的统计口径不同，图上的峰值不要求等于表格数值；筛查后数值仅列于表中。', '图表对应')
        snrs = [(c, metric(c,'ROI','snr_db','analyzable',1).get('mean')) for c in conditions]
        p('各注视条件基频 SNR 均值为：' + '；'.join(f'{c} Hz：{shown(v)} dB' for c,v in snrs) + '。高于邻近背景可作为频率响应线索，还需结合静息对照与通道分布，不能单独判定设备质量。', '结果解释')
        if evidence.get('harmonic_overlaps'):
            p('存在目标间共享谐波频率，不能将这些频率的响应分别归属于单一目标。')
        doc.add_heading('与独立睁眼静息的对照', 2)
        baseline=[r for r in evidence.get('baseline_comparison',[]) if r['channel']=='ROI' and r['subset']=='analyzable']
        if baseline:
            table(['目标 Hz', '任务功率', '睁眼功率', '变化 dB', '任务 n／基线窗 n'], [
                (r['condition'], shown(r['task_power_uv2']), shown(r['baseline_power_uv2']), shown(r['change_db']),
                 f"{r['n_trials']}／{r['n_baseline_windows']}") for r in baseline], [.8,1.2,1.2,1,2.3])
            p('功率单位 μV²，变化为 10log10(任务平均功率/睁眼平均功率)；匹配电极、窗长与频谱方法。视野和注视未独立验证，仅作描述性对照。')
        else:
            p(evidence.get('resting_baseline',{}).get('reason','缺少匹配睁眼基线'),'静息对照不可用')
        doc.add_heading('逐通道响应分布', 2)
        figure_named('response-channels.png')
        p('每格为对应目标条件、对应电极的逐试次基频 SNR 均值，仅纳入可分析试次；上方 SNR 表汇总的是分析电极 ROI。通道值与 ROI 值不能直接逐值对应，通道统计与样本量见 response_summary.csv。', '通道解释')
    else:
        blo,bhi=report['configuration']['mi_baseline_relative_fixation_sec']
        lo,hi=report['configuration']['mi_task_relative_imagery_sec']
        p(f'试次准备后 {blo:g}–{bhi:g} 秒为基线，想象后 {lo:g}–{hi:g} 秒为任务窗。').paragraph_format.keep_with_next = True
        p('ERD=100×(任务功率−试次准备基线功率)/基线功率；μ 8–13 Hz，β 13–30 Hz；负值为功率下降。独立静息不替代试次基线。').paragraph_format.keep_with_next = True
        start = group_start()
        doc.add_heading('基线与任务功率', 2)
        table(['条件／通道', 'μ 基线／任务功率', 'β 基线／任务功率', '可分析／筛查后 n'], [
            (f'{c} {ch}', '／'.join(shown(metric(c,ch,k).get('mean')) for k in ('mu_baseline_power_uv2','mu_task_power_uv2')),
             '／'.join(shown(metric(c,ch,k).get('mean')) for k in ('beta_baseline_power_uv2','beta_task_power_uv2')),
             f"{metric(c,ch,'mu_erd_pct').get('n',0)}／{metric(c,ch,'mu_erd_pct','screened').get('n',0)}")
            for c in conditions for ch in ('C3','C4')], [1.1,1.8,1.8,1.8])
        p('功率单位 μV²；n 为对应条件／通道的有效试次数，β 和各扩展指标的具体 n 另存明细。此表用于核对 ERD 的基线分母。')
        figure_named('response-mi-psd.png')
        p('表用完整基线／任务分析窗计算频带积分功率均值；图用同批试次的 1 秒 Welch 分段谱（50% 重叠），虚线为基线、实线为任务。单位及谱估计窗长不同，图用于观察频谱形状，不能直接从曲线重建表中数值。', '功率解释')
        group_end(start)
        start = group_start()
        doc.add_heading('μ／β 去同步及时间变化', 2)
        table(['条件／通道','μ ERD % [95%区间]','β ERD % [95%区间]','筛查后 μ／β'], [
            (f'{c} {ch}',estimate(metric(c,ch,'mu_erd_pct')),estimate(metric(c,ch,'beta_erd_pct')),
             '／'.join(shown(metric(c,ch,k,'screened').get('mean')) for k in ('mu_erd_pct','beta_erd_pct'))) for c in conditions for ch in ('C3','C4')], [1.1,1.85,1.85,1.7])
        figure_named('mi-erd.png')
        curves = task.get('timecourses', [])
        same = bool(curves) and all(r.get('clean') for r in curves)
        enough = any(len({r['trial'] for r in curves if r['condition'] == c}) >= 3 for c in ('左手','右手','静息'))
        p('表为完整任务分析窗的逐试次 ERD 均值，曲线为各滑动时间窗的 ERD 均值。'
          + ('本次可分析与筛查后试次相同，合并显示一排曲线。' if same else '上下排分别对应可分析与筛查后试次。')
          + ('仅试次数足够的条件绘制 95% 自助法区间阴影；不足时不绘制。' if enough else '各条件不足 3 次，不估计置信区间，不绘制阴影。')
          + '不同时间窗的曲线点不应与整窗表格逐值比较。', '时程解释')
        p('负值表示相对本试次准备基线功率下降，正值表示上升；需同时观察左右手、C3／C4 和静息条件。单次未出现预期下降不能独立归因于设备。')
        group_end(start)
        doc.add_heading('响应的时频分布', 2)
        figure_named('response-mi-tfr.png')
        p('图显示可分析试次在各频率、各时窗相对自身准备基线的功率变化，单位为 dB；与上节 ERD（%）共享试次基线，但频率分辨率、时间窗和数值尺度不同。蓝色表示下降；色阶固定 ±10 dB，超出范围仍按端点颜色显示。', '时频解释')
        start = group_start()
        doc.add_heading('左右侧化', 2)
        table(['条件', 'μ C3−C4（dB）', 'μ 对侧−同侧（dB）', 'β 对侧−同侧（dB）'], [
            (c, estimate(metric(c,'C3-C4','mu_lateral_db')),
             estimate(metric(c,'contra-ipsi','mu_change_db')), estimate(metric(c,'contra-ipsi','beta_change_db')))
            for c in conditions], [1,1.9,1.8,1.8])
        figure_named('response-mi-lateral.png')
        p('图为单试次 μ 功率变化的 C3−C4 差，表中第二列为这些点的均值及 95% 区间。对侧−同侧列使用另一符号约定：左手为 C4−C3，右手为 C3−C4；因此左手条件两列符号相反，右手相同。静息不定义对侧。', '侧化对应')
        p('对侧−同侧为负时表示对侧相对抑制较强；图中圆点表示幅值筛查保留，叉号表示标记。±0.6 dB 是固定分类规则，不是设备合格线。')
        group_end(start)
        doc.add_heading('休息早期的 β 恢复', 2)
        recovery = [r for r in summary if r['metric']=='beta_recovery_erd_pct' and r['subset']=='analyzable']
        if recovery:
            table(['恢复期条件／通道', 'n', 'β 变化 % [95%区间]'], [
                (f"{r['condition']} {r['channel']}",r['n'],estimate(r)) for r in recovery], [2.1,.5,3.9])
            p('恢复窗为休息开始后 0.5–1.5 秒，相对同一试次准备基线。仅描述早期恢复，不保证覆盖 β 反弹峰值；未见反弹不等于设备差。')
        else:
            p('没有足长且有效的恢复窗，恢复期 β 变化不可用。')
    if summary:
        p('响应区间为试次自助法描述，不足 3 次不估计；同一被试内试次可能相关。筛查前后比较用于观察异常试次的影响。')
        if len(task.get('blocks',{})) > 1:
            doc.add_heading('分组响应稳定性', 2)
            for label,path in figures:
                if path.name == 'response-blocks.png':
                    layout.figure(label,path)
            p('观察响应随实验分组的变化；图中为可分析试次均值，各组 n、离散度与区间见 block_response.csv，不代表跨会话重测信度。')

    start = group_start()
    doc.add_heading('固定算法与综合结论', 1)
    if task['status']=='available':
        p(f"固定 FBCCA，三个子带与三次谐波，不训练个人模板；{len(run.get('protocol_start_payload', {}).get('ssvep_frequencies', [])) or 4} 目标，名义随机水平 {100/(len(run.get('protocol_start_payload', {}).get('ssvep_frequencies', [])) or 4):g}%。" if ssvep else
          '固定 μ 频带 C3−C4 功率变化 dB 差 >0.6 判左手、<−0.6 判右手，其余未判定；静息不计入二分类，名义随机水平 50%。')
        table(['范围','n','准确率 [95%区间]','平衡准确率','判定覆盖率'], [
            (label,task[key]['n'],shown(task[key]['accuracy'],True)+' ['+('～'.join(shown(v,True) for v in task[key]['ci95_wilson']) if task[key]['ci95_wilson'] else '不可用')+']',
             shown(task[key]['balanced_accuracy'],True),shown(task[key]['coverage'],True))
            for label,key in [('全部已观察','all_trials'),('可分析','analyzable_trials'),('幅值筛查后','clean_trials')]], [1.1,.4,2.3,1.4,1.3])
        overall=task['all_trials']
        p(f"全部已观察计分试次中，数据不可分析 {overall['data_unavailable']} 次，模型不确定 {overall['model_abstained']} 次。")
        p('未判定在总体准确率中按未正确处理；未观察到的计划试次不纳入分类分母。分类只作任务信息可分性的辅助证据，不等于行为正确率或设备总体质量。')
    else:
        p(task.get('reason','任务证据不足'),'分类暂不可用')
    for label,path in figures:
        if path.name == 'task-confusion.png':
            layout.figure(label,path,width=5.2)
            p('矩阵对应上表“全部已观察”一行，不对应筛查后子集。每格同时标出试次数和该真实条件内的比例；“模型不确定”和“数据不可分析”均计入分母。', '分类对应')
    group_end(start)
    doc.add_heading('最终综合结论', 2)
    flagged=[c['channel'] for c in channels if (c.get('flat_fraction') or 0)>0 or (c.get('large_amplitude_fraction') or 0)>0 or c['finite_fraction']<1]
    p('仅验证流程，不能据模拟信号评价设备。' if simulated else
      (f'发现 {len(flagged)} 个通道存在幅值或非有限值线索，需复核：'+ '、'.join(flagged[:12])+('等。' if len(flagged)>12 else '。') if flagged else ('已执行的幅值筛查未标记明显异常；不排除其他干扰。' if any(c.get('flat_fraction') is not None for c in channels) else '缺少已确认单位与类型的通道，幅值质量暂不可评估。')), '信号质量')
    p('响应可作本次描述，需结合电极位置、任务执行和同步证据解释。' if task['status']=='available' else '任务证据不足，不能据此判断设备表现差。','任务响应')
    p('接触结论仅限本次测量时点。' if report['impedance'] else '缺少实际接触测量，不能评价。','电极接触')
    p('优先核对异常窗口、排除原因、参考方式和同步依据，再按相同条件复测；单次结果不能支持跨设备排名、能力判断或医疗诊断。','建议')
    if evidence.get('errors'):
        p(f"{len(evidence['errors'])} 个试次补充响应计算失败，原因见 response_errors.csv。")
    if report.get('general_features',{}).get('errors'):
        p('部分通用指标计算不可用，原因见 general_errors.csv，不以零替代。')
    # Keep condition and short block tables together; long channel lists may span pages.
    for result_table in doc.tables:
        if len(result_table.rows) <= 7:
            for index, row in enumerate(result_table.rows):
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.keep_with_next = index < len(result_table.rows) - 1
    # Keep a small evidence topic together, including the table, plot and explanation.
    for group in paired_groups:
        for element in group[:-1]:
            if element.tag.endswith('}p'):
                Paragraph(element, doc).paragraph_format.keep_with_next = True
            elif element.tag.endswith('}tbl'):
                for row in Table(element, doc).rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            paragraph.paragraph_format.keep_with_next = True
    doc.save(destination)


def compact_figure(report, destination):
    """One readable two-panel response overview; detailed figures remain exported."""
    from matplotlib import pyplot as plt
    ssvep = report['run']['paradigm'] == 'SSVEP'
    spectra = report['task'].get('response_evidence', {}).get('spectra', [])
    curves = report['task'].get('timecourses', [])
    if not spectra and not curves:
        return []
    fig, axes = plt.subplots(1, 2, figsize=(7, 2.7), layout='constrained')
    if ssvep:
        chosen = [r for r in spectra if r['subset']=='analyzable' and r['channel']=='ROI']
        for row in chosen:
            f=np.array(row['frequency_hz'])
            line,=axes[0].semilogy(f,np.maximum(row['psd'],1e-20),label=f"{row['condition']} Hz")
            axes[1].plot(f,np.array(row['snr_db'],dtype=float),color=line.get_color())
        for ax in axes:
            ax.set(xlim=(4,min(report['device']['sfreq']/2,max([65]+[float(r['condition'])*3+2 for r in chosen]))),xlabel='频率 (Hz)')
        axes[0].set(ylabel='PSD (μV²/Hz)',title='条件频谱')
        axes[1].set(ylabel='SNR (dB)',title='局部 SNR 谱')
        axes[0].legend(fontsize=8,ncol=2)
        caption='可分析试次的条件频谱与 SNR 谱；谱图为平均 PSD 的比值，表格为逐试次 dB 均值'
    else:
        for ax,band in zip(axes,('mu','beta')):
            for condition,color in zip(('左手','右手','静息'),('#286a96','#b45d24','#4a8051')):
                for channel,style in (('C3','-'),('C4','--')):
                    selected=[r for r in curves if r['band']==band and r['condition']==condition]
                    times=sorted({r['time_sec'] for r in selected})
                    values=[np.mean([r[channel+'_erd_pct'] for r in selected if r['time_sec']==t]) for t in times]
                    ax.plot(times,values,style,color=color,label=condition+' '+channel)
            ax.axhline(0,color='gray',linewidth=.6)
            ax.set(xlabel='想象起点后 (s)',ylabel='ERD (%)',title=band+' 全部可分析')
        axes[0].legend(fontsize=7,ncol=2)
        caption='可分析试次的 μ／β ERD 均值时程；负值为功率下降，筛查前后数值见表'
    for ax in axes:
        ax.grid(alpha=.2)
    path=destination/'compact-response.png'
    fig.savefig(path,dpi=180)
    plt.close(fig)
    return [(caption,path)]
